#!/usr/bin/env python3
"""Genera los datos del prototipo (prototipo/data/*.js) a partir de fuentes abiertas de Luxemburgo.

Dependencias: build/requirements.txt. Antes: python3 prototipo/build/get_gtfs.py. Uso: python3 prototipo/build/build_data.py

Descargas en ~/.cache/nether_app/raw (fuera del repositorio: son grandes); las manuales, en prototipo/README.md. Fuentes y licencias: notas/fuentes_datos.md.
"""
import csv, datetime, json, os, sqlite3, sys, time, unicodedata, urllib.parse, urllib.request
import numpy as np
import re
import openpyxl, shapefile, xlrd
from pyproj import Transformer
from shapely import wkb, points
from shapely.geometry import shape, LineString, Polygon
from shapely.strtree import STRtree

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.normpath(os.path.join(HERE, '..', 'data'))
RAW = os.path.expanduser('~/.cache/nether_app/raw')
GTFS = os.path.expanduser('~/.cache/nether_gtfs')   # caché de horarios compartida con otras herramientas locales
REGION = (49.55, 6.03, 49.70, 6.28)                  # lat0, lon0, lat1, lon1
STEP_KM = 0.15                                        # lado de la celda
KX, KY = 72.1, 111.2                                  # km por grado de lon / lat a esta latitud
UA = {'User-Agent': 'nether-app-prototype/0.1'}

URLS = {
    'quartiers.geojson': 'https://maps.vdl.lu/arcgis/rest/services/OPENDATA/GEOJSON/FeatureServer/24/query?where=1%3D1&outFields=*&outSR=4326&f=geojson',
    'communes.geojson': 'https://download.data.public.lu/resources/limites-administratives-du-grand-duche-de-luxembourg/20231123-101528/communes4326.geojson',
    'loyers_quartiers.xlsx': 'https://download.data.public.lu/resources/loyers-annonces-des-logements-a-luxembourg-ville-par-quartier/20260924-085714/vdl-location-appartements-2025-26.xlsx',
    'loyers_communes.xls': 'https://download.data.public.lu/resources/loyers-annonces-des-logements-par-commune/20260924-085311/location-appartement-2025-26.xls',
    'agglo.gpkg': 'https://download.data.public.lu/resources/environmental-noise-strategic-noise-maps/20250724-133722/lu-agglomeration-2021-strategicnoisemaps.gpkg',
    'airport.gpkg': 'https://download.data.public.lu/resources/environmental-noise-strategic-noise-maps/20250724-133721/lu-majorairports-2021-strategicnoisemaps.gpkg',
    'majorroads.gpkg': 'https://download.data.public.lu/resources/environmental-noise-strategic-noise-maps/20250724-133726/lu-majorroads-2021-strategicnoisemaps.gpkg',
    'majorrailways.gpkg': 'https://download.data.public.lu/resources/environmental-noise-strategic-noise-maps/20250724-133723/lu-majorrailways-2021-strategicnoisemaps.gpkg',
    'aggsource.gpkg': 'https://download.data.public.lu/resources/environmental-noise-report-on-all-major-roads-railways-airports-and-agglomerations-df1-and-df5/20250724-134505/agglomerationsource-2020-lu.gpkg',
}

POI_QUERIES = {   # categoría: filtros de OpenStreetMap
    'supermarket': ['nwr["shop"="supermarket"]'],
    'pharmacy': ['nwr["amenity"="pharmacy"]'],
    'hospital': ['nwr["amenity"="hospital"]'],
    'doctor': ['nwr["amenity"~"^(doctors|clinic)$"]', 'nwr["healthcare"="doctor"]'],
    'school': ['nwr["amenity"="school"]'],
    'kindergarten': ['nwr["amenity"~"^(kindergarten|childcare)$"]'],
    'university': ['nwr["amenity"~"^(university|college)$"]'],
    'park': ['nwr["leisure"="park"]'],
    'gym': ['nwr["leisure"="fitness_centre"]'],
    'nightlife': ['nwr["amenity"~"^(bar|pub|nightclub)$"]'],
}

to_luref = Transformer.from_crs(4326, 2169, always_xy=True)   # metros, Luxemburgo
to_3035 = Transformer.from_crs(4326, 3035, always_xy=True)    # proyección de los mapas de ruido


def log(*a): print(*a, file=sys.stderr, flush=True)


def raw(name):
    path = os.path.join(RAW, name)
    if not os.path.exists(path):
        os.makedirs(RAW, exist_ok=True)
        log('descargando', name)
        req = urllib.request.Request(URLS[name], headers=UA)
        with urllib.request.urlopen(req) as r, open(path, 'wb') as f:
            while True:
                b = r.read(1 << 20)
                if not b: break
                f.write(b)
    return path


def norm(s):
    s = unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode().lower()
    return s.replace('-', ' ').replace('/', ' ').split()[0] if s.strip() else ''


# ---------- rejilla ----------
def grid():
    lat0, lon0, lat1, lon1 = REGION
    dlat, dlon = STEP_KM / KY, STEP_KM / KX
    rows, cols = int((lat1 - lat0) / dlat), int((lon1 - lon0) / dlon)
    la = lat0 + (np.arange(rows) + 0.5) * dlat
    lo = lon0 + (np.arange(cols) + 0.5) * dlon
    LO, LA = np.meshgrid(lo, la)          # fila 0 = sur
    return dict(lat0=lat0, lon0=lon0, dlat=dlat, dlon=dlon, rows=rows, cols=cols), LA.ravel(), LO.ravel()


# ---------- ruido ----------
BAND = {'Lden5559': 55, 'Lden6064': 60, 'Lden6569': 65, 'Lden7074': 70, 'LdenGreaterThan75': 75}


def gpkg_geom(blob):
    b = bytes(blob)
    env = (b[3] >> 1) & 7
    return wkb.loads(b[8 + {0: 0, 1: 32, 2: 48, 3: 48, 4: 64}[env]:])


def gpkg_layer(path, table):
    con = sqlite3.connect(path)
    out = [(gpkg_geom(blob), BAND.get(cat, 0)) for blob, cat in con.execute(f'select geom, category from "{table}"')]
    con.close()
    return out


def noise_coverage(LA, LO):
    """2 = dentro de la aglomeración AggloLux (todas las calles, tren y aeropuerto); 1 = fuera (solo grandes ejes y aeropuerto)."""
    con = sqlite3.connect(raw('aggsource.gpkg'))
    blob, = con.execute("select geom from AgglomerationSource where agglomerationId_identifier = 'AG_LU_00_1'").fetchone()
    con.close()
    x, y = to_3035.transform(LO, LA)
    inside = gpkg_geom(blob).contains(points(x, y))
    return np.where(inside, 2, 1).astype(np.int8)


def noise(LA, LO, layers):
    x, y = to_3035.transform(LO, LA)
    pts = points(x, y)
    res = np.zeros(len(LA), dtype=np.int16)
    for path, table in layers:
        feats = gpkg_layer(path, table)
        if not feats: continue
        tree = STRtree([g for g, _ in feats])
        bands = np.array([b for _, b in feats], dtype=np.int16)
        ip, ig = tree.query(pts, predicate='intersects')
        np.maximum.at(res, ip, bands[ig])
        log('  ruido', table, len(feats), 'polígonos')
    return res


AIRPORT_2023 = os.path.join(RAW, 'airport2023', 'snm_MajorAirport_2023_{}')   # del zip de documentación 2023 (CC0)


def airport_2023(LA, LO, kind):
    """Mapa del aeropuerto 2023 (actualización intermedia, Administration de l'environnement, 2025).
    Shapefile en LUREF (EPSG:2169) aunque el .prj diga Web Mercator. Banda = LEVEL_1; >= 80 cuenta como 75."""
    r = shapefile.Reader(AIRPORT_2023.format(kind))
    geoms = [shape(sh.__geo_interface__).buffer(0) for sh in r.shapes()]
    bands = np.array([min(int(rec['LEVEL_1']), 75) for rec in r.records()], dtype=np.int16)
    x, y = to_luref.transform(LO, LA)
    ip, ig = STRtree(geoms).query(points(x, y), predicate='intersects')
    res = np.zeros(len(LA), dtype=np.int16)
    np.maximum.at(res, ip, bands[ig])
    log('  ruido aeropuerto 2023', kind, len(geoms), 'polígonos')
    return res


def elevation(LA, LO):
    """Altitud (m) por celda: punto más cercano del MNT 75 m (ACT, CC0; fichero luxaltblh50 = lat, lon, altura)."""
    path = os.path.join(RAW, 'mnt75', 'luxaltblh50')
    lat0, lon0, lat1, lon1 = REGION
    d = np.loadtxt(path)
    d = d[(d[:, 0] > lat0 - 0.01) & (d[:, 0] < lat1 + 0.01) & (d[:, 1] > lon0 - 0.01) & (d[:, 1] < lon1 + 0.01)]
    px, py = to_luref.transform(d[:, 1], d[:, 0])
    cx, cy = to_luref.transform(LO, LA)
    idx = STRtree(points(px, py)).query_nearest(points(cx, cy), all_matches=False)[1]
    z = np.round(d[idx, 2]).astype(int)
    # En la ciudad hay un MNT de 4 m en el mismo zip (test.dem, ESRI ASCII en LUREF): mediana dentro de la celda.
    dem = os.path.join(RAW, 'mnt75', 'test.dem')
    if os.path.exists(dem):
        with open(dem) as f:
            h = {k: float(v) for k, v in (f.readline().split() for _ in range(5))}
            Z = np.loadtxt(f)
        nr, cs = int(h['nrows']), h['cellsize']
        half = STEP_KM * 1000 / 2
        n4 = 0
        for i in range(len(LA)):
            c0 = int((cx[i] - half - h['xllcenter']) / cs); c1 = int((cx[i] + half - h['xllcenter']) / cs)
            r1 = nr - 1 - int((cy[i] - half - h['yllcenter']) / cs); r0 = nr - 1 - int((cy[i] + half - h['yllcenter']) / cs)
            if c0 >= 0 and r0 >= 0 and c1 < Z.shape[1] and r1 < Z.shape[0]:
                z[i] = int(np.median(Z[r0:r1 + 1, c0:c1 + 1])); n4 += 1
        log('  altitud: MNT 4 m en', n4, 'celdas; MNT 75 m en el resto')
    return z


# ---------- barrios, municipios y alquiler ----------
def rents():
    wb = openpyxl.load_workbook(raw('loyers_quartiers.xlsx'), data_only=True)
    q = {}
    for row in wb.worksheets[0].iter_rows(values_only=True):
        v = [x for x in row if x is not None]
        if len(v) >= 4 and isinstance(v[0], str) and all(isinstance(x, (int, float)) for x in v[1:4]):
            q[norm(v[0])] = (round(float(v[3]), 1), int(v[1]), round(float(v[2])))
    sh = xlrd.open_workbook(raw('loyers_communes.xls')).sheet_by_index(0)
    c = {}
    for i in range(sh.nrows):
        r = sh.row_values(i)
        if isinstance(r[2], str) and isinstance(r[5], float):
            c[norm(r[2])] = (round(r[5], 1), int(r[3]) if isinstance(r[3], float) else 0, round(r[4]) if isinstance(r[4], float) else None)
    return q, c


# Curva oficial €/m² por tamaño (Observatoire de l'Habitat, "Loyers annoncés en 2025 dans l'aire fonctionnelle
# du Luxembourg", 30/09/2026, fig. 5 p. 17; media por clase, asignada al centro de cada clase). Validado por la sesión
# de datos frente a anuncios activos (notas/calidad_datos.md): error absoluto medio ±9 %.
SIZE_CURVE = [(37, 56.3), (60, 31.9), (80, 26.9), (100, 26.1), (120, 22.9), (145, 21.0)]
L_NAT = 0.879   # nivel nacional: 44,19715 €/m² / f(1842,624 / 44,19715)


def f_size(S):
    if S <= SIZE_CURVE[0][0]: return SIZE_CURVE[0][1]
    if S >= SIZE_CURVE[-1][0]: return SIZE_CURVE[-1][1]
    for (s0, v0), (s1, v1) in zip(SIZE_CURVE, SIZE_CURVE[1:]):
        if s0 <= S <= s1:
            t = (np.log(S) - np.log(s0)) / (np.log(s1) - np.log(s0))
            return v0 + t * (v1 - v0)


def level(eur_m2, avg_rent):
    """Nivel de precio de una zona descontando su mezcla de tamaños (1 = curva nacional)."""
    if not eur_m2 or not avg_rent: return None
    return eur_m2 / f_size(avg_rent / eur_m2)


def areas(LA, LO):
    qr, cr = rents()
    names, rent, kind = [], [], []
    geoms, idx = [], []
    for f in json.load(open(raw('quartiers.geojson')))['features']:
        n = f['properties']['FK_QUART_NAME']
        names.append(n); kind.append('district'); rent.append(qr.get(norm(n)))
        geoms.append(shape(f['geometry'])); idx.append(len(names) - 1)
    for f in json.load(open(raw('communes.geojson')))['features']:
        n = f['properties']['COMMUNE']
        if norm(n) == 'luxembourg': continue           # la ciudad va por barrios
        names.append(n); kind.append('commune'); rent.append(cr.get(norm(n)))
        geoms.append(shape(f['geometry'])); idx.append(len(names) - 1)
    tree = STRtree(geoms)
    res = np.full(len(LA), -1, dtype=np.int32)
    ip, ig = tree.query(points(LO, LA), predicate='within')
    res[ip] = np.array(idx)[ig]
    city = cr.get('luxembourg')
    out = []
    for n, k, r in zip(names, kind, rent):
        base = r or (city if k == 'district' else None)
        lv = level(base[0], base[2]) if base else None
        out.append(dict(name=n, kind=k, eur_m2=base[0] if base else None, offers=r[1] if r else 0,
                        avg_rent=base[2] if base else None, estimated=r is None and k == 'district',
                        L=round(lv if lv else L_NAT, 3), L_default=lv is None))
    return out, res


# ---------- puntos de interés (OpenStreetMap) ----------
def overpass(q, cache):
    """Consulta Overpass con reintentos; guarda la respuesta en RAW para no repetirla."""
    path = os.path.join(RAW, cache)
    if os.path.exists(path):
        return json.load(open(path))
    for intento in range(8):
        try:
            data = urllib.parse.urlencode({'data': q}).encode()
            req = urllib.request.Request('https://overpass-api.de/api/interpreter', data=data, headers=UA)
            with urllib.request.urlopen(req, timeout=200) as r:
                res = json.load(r)
            json.dump(res, open(path, 'w'))
            return res
        except Exception as e:
            espera = 20 * (intento + 1)
            log(f'  overpass: {e}; reintento en {espera} s'); time.sleep(espera)
    sys.exit('Overpass no responde; vuelve a ejecutar más tarde (lo ya descargado queda guardado)')


def veloh():
    """Estaciones Vel'oh (GBFS de JCDecaux). Licencia JCDecaux Open Licence: citar "Vel'oh! – JCDecaux"; contactar antes de uso comercial."""
    url = 'https://api.cyclocity.fr/contracts/luxembourg/gbfs/v3/station_information.json'
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=60) as r:
        st = json.load(r)['data']['stations']
    out = []
    for x in st:
        name = x.get('name')
        if isinstance(name, list): name = name[0].get('text', '') if name else ''
        if name and name.startswith('#') and '-' in name: name = name.split('-', 1)[1].title()
        out.append((x['lat'], x['lon'], name or ''))
    return out


# Fusión con Overture Places (regla de la sesión de datos, notas/calidad_datos.md, 2026-10-07)
OVERTURE = {
    'gym': ({'gym', 'sport_or_fitness_facility', 'gymnastics_center'}, 0.8, None, None),
    'kindergarten': ({'child_care_and_day_care', 'day_care_preschool', 'preschool'}, 0.8,
                     r'cr[eè]che|kinder|nursery|maison relais|foyer|spillschoul|garderie', None),
    'supermarket': ({'grocery_store', 'organic_grocery_store', 'discount_store'}, 0.8,
                    r'cactus|delhaize|lidl|aldi|auchan|carrefour|match|colruyt|naturata|alnatura|biogros|massen|pall center|proxy', None),
    'pharmacy': ({'pharmacy'}, 0.75, r'pharmacie|apdikt|apotheke', r'hospilux|medi-market'),
}
NAME_STOP = {'pharmacie', 'creche', 'crèche', 'luxembourg', 'sàrl', 'sarl', 'apotheke', 'apdikt', 'maison', 'relais', 'kinder', 'supermarket', 'supermarché'}


def name_words(n):
    return {w for w in re.findall(r'\w{4,}', (n or '').lower()) if w not in NAME_STOP}


def overture_extra(cat, osm_geoms, osm_names):
    """Lugares de Overture que añaden algo a OSM según la regla; devuelve [(lat, lon, nombre, geom)]."""
    if cat not in OVERTURE: return []
    cats, conf_min, must, never = OVERTURE[cat]
    tree = STRtree(osm_geoms) if osm_geoms else None
    out = []
    for f in json.load(open(os.path.join(RAW, 'overture_places.geojson')))['features']:
        p = f['properties']; tax = (p.get('taxonomy') or {}).get('primary')
        name = ((p.get('names') or {}).get('primary') or '')
        if tax not in cats or (p.get('confidence') or 0) < conf_min: continue
        if must and not re.search(must, name.lower()): continue
        if never and re.search(never, name.lower()): continue
        lo, la = f['geometry']['coordinates'][:2]
        x, y = to_luref.transform(lo, la); g = points(x, y)
        if tree is not None:
            if len(tree.query(g, predicate='dwithin', distance=100)): continue
            near = tree.query(g, predicate='dwithin', distance=300)
            if any(name_words(name) & name_words(osm_names[k]) for k in near): continue
        out.append((la, lo, name, g))
    return out


def sea_childcare(other_geoms, other_shown, kind, cat):
    """Guarderías y maisons relais: base = lista oficial SEA 2021 (Ministère de l'Éducation nationale, CC0),
    geocodificada por la sesión de datos con geoportail.lu; + OSM/Overture a más de 100 m de un SEA."""
    lat0, lon0, lat1, lon1 = REGION
    sea = [x for x in json.load(open(os.path.join(RAW, 'sea_2021_geocoded.json')))
           if x.get('lat') and x.get(kind) and lat0 - 0.03 <= x['lat'] <= lat1 + 0.03 and lon0 - 0.04 <= x['lon'] <= lon1 + 0.04]
    geoms, shown = [], []
    for x in sea:
        px, py = to_luref.transform(x['lon'], x['lat']); geoms.append(points(px, py))
        shown.append([round(x['lat'], 5), round(x['lon'], 5), cat, (x.get('name') or '')[:60]])
    tree = STRtree(geoms)
    kept = 0
    for g, sh in zip(other_geoms, other_shown):
        if not len(tree.query(g, predicate='dwithin', distance=100)):
            geoms.append(g); shown.append([sh[0], sh[1], cat, sh[3]]); kept += 1
    log(f'  {cat}: SEA', len(sea), '+ OSM/Overture', kept)
    return geoms, shown


def pois(LA, LO):
    lat0, lon0, lat1, lon1 = REGION
    bb = f'{lat0 - 0.03},{lon0 - 0.04},{lat1 + 0.03},{lon1 + 0.04}'
    cx, cy = to_luref.transform(LO, LA)
    cells = points(cx, cy)
    dist, shown = {}, []
    for cat, filters in list(POI_QUERIES.items()) + [('bikeshare', None)]:
        geoms, names, start = [], [], len(shown)
        if cat == 'bikeshare':
            for la, lo, name in veloh():
                x, y = to_luref.transform(lo, la); geoms.append(points(x, y)); shown.append([round(la, 5), round(lo, 5), cat, name[:60]])
            els = []
        else:
            q = '[out:json][timeout:120];(' + ''.join(f'{f}({bb});' for f in filters) + ');out geom;'
            els = overpass(q, f'osm_{cat}.json')['elements']
        for e in els:
            name = e.get('tags', {}).get('name', '')
            if e['type'] == 'node':
                x, y = to_luref.transform(e['lon'], e['lat']); g = points(x, y); la, lo = e['lat'], e['lon']
            elif e['type'] == 'way' and e.get('geometry'):
                xy = [to_luref.transform(p['lon'], p['lat']) for p in e['geometry']]
                g = Polygon(xy) if len(xy) > 3 and xy[0] == xy[-1] else LineString(xy) if len(xy) > 1 else None
                if g is None: continue
                la = sum(p['lat'] for p in e['geometry']) / len(e['geometry']); lo = sum(p['lon'] for p in e['geometry']) / len(e['geometry'])
            elif e['type'] == 'relation' and e.get('members'):
                pts = [(p['lon'], p['lat']) for m in e['members'] for p in (m.get('geometry') or [])]
                if not pts: continue
                xs, ys = to_luref.transform([p[0] for p in pts], [p[1] for p in pts])
                g = LineString(list(zip(xs, ys))).convex_hull if cat == 'park' else points(np.mean(xs), np.mean(ys))
                lo, la = np.mean([p[0] for p in pts]), np.mean([p[1] for p in pts])
            else:
                continue
            if cat == 'park' and g.area < 2000: continue      # jardincillos fuera
            if cat == 'school' and re.search(r'musique|music|police|conduite|driving', name.lower()): continue
            geoms.append(g); names.append(name)
            shown.append([round(float(la), 5), round(float(lo), 5), cat, name[:60]])
        extra = overture_extra(cat, geoms, names)
        for la, lo, name, g in extra:
            geoms.append(g); names.append(name)
            shown.append([round(float(la), 5), round(float(lo), 5), cat, name[:60]])
        if extra: log('  + Overture', cat, len(extra))
        if cat == 'kindergarten':   # se parte en guarderías (0-4, + OSM/Overture) y maisons relais (escolares)
            ng, ns = sea_childcare(geoms, shown[start:], 'young_children', 'nursery')
            ag, as_ = sea_childcare([], [], 'school_children', 'afterschool')
            shown[start:] = ns + as_
            for c, gs in (('nursery', ng), ('afterschool', ag)):
                _, d = STRtree(gs).query_nearest(cells, return_distance=True, all_matches=False)
                dist[c] = np.minimum(np.round(d / 10), 600).astype(int)
            continue
        tree = STRtree(geoms)
        _, d = tree.query_nearest(cells, return_distance=True, all_matches=False)
        dist[cat] = np.minimum(np.round(d / 10), 600).astype(int)   # decenas de metros, tope 6 km
        log('  POI', cat, len(geoms))
    return dist, shown


# ---------- transporte público (GTFS) ----------
def secs(s): h, m, x = s.split(':'); return int(h) * 3600 + int(m) * 60 + int(x)


def transit():
    hoy = datetime.date.today()
    dia = hoy + datetime.timedelta(days=(1 - hoy.weekday()) % 7 or 7)   # próximo martes
    D = dia.strftime('%Y%m%d')
    act = set()
    for r in csv.DictReader(open(f'{GTFS}/calendar.txt', encoding='utf-8-sig')):
        if r['tuesday'] == '1' and r['start_date'] <= D <= r['end_date']: act.add(r['service_id'])
    for r in csv.DictReader(open(f'{GTFS}/calendar_dates.txt', encoding='utf-8-sig')):
        if r['date'] == D: (act.add if r['exception_type'] == '1' else act.discard)(r['service_id'])
    trips = {r['trip_id'] for r in csv.DictReader(open(f'{GTFS}/trips.txt', encoding='utf-8-sig')) if r['service_id'] in act}
    lat0, lon0, lat1, lon1 = REGION
    stops = {}
    for r in csv.DictReader(open(f'{GTFS}/stops.txt', encoding='utf-8-sig')):
        la, lo = float(r['stop_lat']), float(r['stop_lon'])
        if lat0 - 0.03 <= la <= lat1 + 0.03 and lon0 - 0.04 <= lo <= lon1 + 0.04: stops[r['stop_id']] = (la, lo)
    sidx, tidx, conns, prev = {}, {}, [], {}
    for r in csv.DictReader(open(f'{GTFS}/stop_times.txt', encoding='utf-8-sig')):
        t = r['trip_id']
        if t not in trips: continue
        cur = (r['stop_id'], secs(r['arrival_time']), secs(r['departure_time']), int(r['stop_sequence']),
               r.get('pickup_type') or '0', r.get('drop_off_type') or '0')
        p = prev.get(t); prev[t] = cur
        if p and p[3] < cur[3] and p[0] in stops and cur[0] in stops and 6 * 3600 <= p[2] <= 10.5 * 3600:
            a = sidx.setdefault(p[0], len(sidx)); b = sidx.setdefault(cur[0], len(sidx)); k = tidx.setdefault(t, len(tidx))
            flags = (1 if p[4] != '1' else 0) | (2 if cur[5] != '1' else 0)   # 1 = se puede subir en a, 2 = se puede bajar en b
            conns.append((p[2], cur[1], a, b, k, flags))
    conns.sort(key=lambda c: -c[0])
    st = [None] * len(sidx)
    for s, i in sidx.items(): st[i] = [round(stops[s][0], 5), round(stops[s][1], 5)]
    log('  transporte', dia, len(st), 'paradas', len(conns), 'conexiones')
    return dict(day=str(dia), stride=6, stops=st, conns=[v for c in conns for v in c])


# Fuentes que no están en URLS (descargadas aparte) para el manifiesto
EXTRA_SOURCES = {
    'airport2023': ('https://download.data.public.lu/resources/environmental-noise-strategic-noise-maps/20250327-151652/strategic-noise-maps-2023-on-major-airport-documentation.zip', 'airport2023/snm_MajorAirport_2023_Lden.shp'),
    'mnt75': ('https://download.data.public.lu/resources/modele-numerique-de-terrain-75m/20160405-133110/mnt_75.zip', 'mnt75/luxaltblh50'),
    'sea_2021': ('https://download.data.public.lu/resources/adresses-des-batiments-scolaires-2021/20230328-230303/adr-sea-2021.xlsx', 'sea_2021_geocoded.json'),
    'overture_places': ('overturemaps download --bbox=6.00,49.52,6.33,49.73 -f geojson --type=place', 'overture_places.geojson'),
}


def dpl_last_modified(url):
    """last_modified del recurso en data.public.lu (None si no es de allí o no responde)."""
    if 'download.data.public.lu/resources/' not in url: return None
    slug = url.split('/resources/')[1].split('/')[0]
    try:
        req = urllib.request.Request(f'https://data.public.lu/api/1/datasets/{slug}/', headers=UA)
        with urllib.request.urlopen(req, timeout=30) as r:
            for res in json.load(r)['resources']:
                if res['url'] == url: return res.get('last_modified')
    except Exception as e:
        log('  manifiesto: sin respuesta de data.public.lu para', slug, e)
    return None


def manifest(tr, shown, area_list, built):
    """prototipo/data/manifest.json: versiones de cada fuente y recuentos, para comparar mes a mes (no se publica)."""
    def file_info(path):
        return dict(local=path.replace(os.path.expanduser('~'), '~'), size=os.path.getsize(path),
                    downloaded=datetime.datetime.fromtimestamp(os.path.getmtime(path)).isoformat(timespec='seconds')) if os.path.exists(path) else None
    src = {}
    for name, url in URLS.items():
        src[name] = dict(url=url, last_modified=dpl_last_modified(url), **(file_info(os.path.join(RAW, name)) or {}))
    for name, (url, rel) in EXTRA_SOURCES.items():
        src[name] = dict(url=url, last_modified=dpl_last_modified(url), **(file_info(os.path.join(RAW, rel)) or {}))
    try:
        ov = json.load(open(os.path.join(RAW, 'overture_places.geojson')))['features']
        rel = sorted({s_.get('version') for f in ov for s_ in f['properties'].get('sources') or [] if s_.get('dataset') == 'Overture'})
        src['overture_places'].update(release=rel[-1] if rel else None, features=len(ov))
    except Exception:
        pass
    cal = os.path.join(GTFS, 'calendar.txt')
    dates = [(r['start_date'], r['end_date']) for r in csv.DictReader(open(cal, encoding='utf-8-sig'))] if os.path.exists(cal) else []
    src['gtfs'] = dict(url='https://data.public.lu/en/datasets/horaires-et-arrets-des-transport-publics-gtfs/',
                       valid_from=min(d[0] for d in dates) if dates else None, valid_to=max(d[1] for d in dates) if dates else None,
                       modelled_day=tr['day'], stops=len(tr['stops']), connections=len(tr['conns']) // tr.get('stride', 5),
                       **(file_info(os.path.join(GTFS, 'stop_times.txt')) or {}))
    src['veloh'] = dict(url='https://api.cyclocity.fr/contracts/luxembourg/gbfs/v3/station_information.json')
    counts = {}
    for _, _, cat, _ in shown: counts[cat] = counts.get(cat, 0) + 1
    m = dict(built=built, region=REGION, step_km=STEP_KM, sources=src, poi_counts=counts,
             areas=dict(n=len(area_list), with_rent=sum(1 for a in area_list if a['eur_m2']),
                        L={a['name']: a['L'] for a in area_list if not a['L_default']}))
    with open(os.path.join(OUT, 'manifest.json'), 'w') as f:
        json.dump(m, f, indent=1, ensure_ascii=False)
    log('escrito manifest.json')


def main():
    os.makedirs(OUT, exist_ok=True)
    g, LA, LO = grid()
    log('rejilla', g['rows'], 'x', g['cols'], '=', len(LA), 'celdas')
    log('ruido')
    road = noise(LA, LO, [(raw('agglo.gpkg'), 'NoiseContours_roadsInAgglomeration_Lden'),
                          (raw('agglo.gpkg'), 'NoiseContours_majorroadsInAgglomeration_Lden'),
                          (raw('majorroads.gpkg'), 'NoiseContours_majorRoadsIncludingAgglomeration_Lden')])
    rail = noise(LA, LO, [(raw('agglo.gpkg'), 'NoiseContours_railwaysInAgglomeration_Lden'),
                          (raw('agglo.gpkg'), 'NoiseContours_majorrailwaysInAgglomeration_Lden'),
                          (raw('majorrailways.gpkg'), 'NoiseContours_majorRailwaysIncludingAgglomeration_Lden')])
    air = airport_2023(LA, LO, 'Lden')          # 2021 sobrestimaba cerca de la pista (notas/calidad_datos.md)
    air_night = airport_2023(LA, LO, 'Lngt')    # desde 45 dB
    elev = elevation(LA, LO)
    cov = noise_coverage(LA, LO)
    log('  cobertura AggloLux:', int((cov == 2).sum()), 'de', len(cov), 'celdas')
    log('barrios y alquiler')
    area_list, area = areas(LA, LO)
    log('puntos de interés')
    dist, shown = pois(LA, LO)
    log('transporte')
    tr = transit()
    built = datetime.date.today().isoformat()
    grid_js = dict(g, built=built, area=area.tolist(), noise_road=road.tolist(), noise_rail=rail.tolist(), noise_air=air.tolist(), noise_air_night=air_night.tolist(), noise_cov=cov.tolist(), elev=elev.tolist(),
                   dist={k: v.tolist() for k, v in dist.items()})
    for name, var, obj in (('grid.js', 'GRID', grid_js), ('areas.js', 'AREAS', area_list),
                           ('pois.js', 'POIS', shown), ('transit.js', 'TRANSIT', tr)):
        with open(os.path.join(OUT, name), 'w') as f:
            f.write(f'window.{var}=' + json.dumps(obj, separators=(',', ':')) + ';\n')
        log('escrito', name, os.path.getsize(os.path.join(OUT, name)) // 1024, 'KB')
    manifest(tr, shown, area_list, built)


if __name__ == '__main__':
    main()
